"""CPU-only API regression on controlled scenarios; never a model-quality benchmark."""
import argparse
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def material(text, fmt):
    if fmt == "txt":
        return text.encode()
    buf = io.BytesIO()
    if fmt == "docx":
        from docx import Document
        doc = Document()
        for line in text.splitlines():
            doc.add_paragraph(line)
        doc.save(buf)
    elif fmt == "xlsx":
        from openpyxl import Workbook
        book = Workbook()
        for index, line in enumerate(text.splitlines(), 1):
            book.active.cell(index, 1, line)
        book.save(buf)
    elif fmt == "pdf":
        from reportlab.pdfgen.canvas import Canvas
        page = Canvas(buf)
        cursor = page.beginText(40, 800)
        for line in text.splitlines():
            cursor.textLine(line)
        page.drawText(cursor)
        page.save()
    else:
        raise ValueError(fmt)
    return buf.getvalue()


def evaluate(path, output):
    # Environment changes are confined to this standalone benchmark process.
    os.environ.pop("TRADEPILOT_MODEL_URL", None)
    from fastapi.testclient import TestClient
    from tradepilot.api import create_app
    from tradepilot.engine import evidence_valid
    source = path.read_bytes()
    dataset = json.loads(source)
    results = []
    with tempfile.TemporaryDirectory(prefix="tradepilot-benchmark-") as temp:
        with TestClient(create_app(data_dir=temp, token="regression-only"), headers={"X-API-Key": "regression-only"}) as client:
            for spec in dataset["scenarios"]:
                started = time.monotonic()
                result = {"id": spec["id"], "category": spec["category"], "passed": False}
                try:
                    def check(response):
                        response.raise_for_status()
                        return response.json()
                    body = {"title": spec["id"]}
                    if spec.get("policy"):
                        body["policy"] = spec["policy"]
                    case = check(client.post("/cases", json=body))
                    prefix = f"/cases/{case['id']}"
                    for index, doc in enumerate(spec["documents"]):
                        check(client.post(prefix + "/documents", data={"kind": doc["kind"], "reference": str(doc["reference"]).lower()},
                            files={"file": (f"{index}.{doc['format']}", material(doc["text"], doc["format"]))}))
                    run = check(client.post(prefix + "/runs"))
                    endpoint = prefix + "/runs/" + run["id"]
                    run = check(client.get(endpoint))
                    assert run["status"] == "awaiting_review", run.get("error")
                    actual = {(f["kind"], f["title"]) for f in run["findings"]}
                    expected = {(f["kind"], f["title"]) for f in spec["expected"]}
                    result["missing"] = sorted(expected - actual)
                    result["unexpected"] = sorted(actual - expected)
                    assert actual == expected, "审查结果与业务规则预期不一致"
                    assert all(evidence_valid(e, run["documents"]) for f in run["findings"] for e in f["evidence"])
                    assert client.post(endpoint + "/approve", json={}).status_code == 409
                    for f in run["findings"]:
                        check(client.post(endpoint + "/decisions/" + f["id"], json={"action":"accept", "reason":"工程回归中显式人工确认动作"}))
                    approved = check(client.post(endpoint + "/approve", json={"acknowledge_limitations":True, "reason":"确认仅规则模式覆盖范围"}))
                    assert approved["status"] == "approved"
                    assert check(client.post(endpoint + "/approve", json={})) ["status"] == "approved"
                    assert check(client.get(prefix))["status"] == "archived"
                    exported = check(client.get(endpoint + "/export/json"))
                    assert exported["run"]["findings"] == run["findings"]
                    result["findings"] = len(run["findings"])
                    result["passed"] = True
                except Exception as exc:
                    result["error"] = f"{type(exc).__name__}: {str(exc)[:600]}"
                result["seconds"] = round(time.monotonic() - started, 3)
                results.append(result)
    passed = sum(r["passed"] for r in results)
    report = {"benchmark": dataset["version"], "input_sha256": hashlib.sha256(source).hexdigest(),
              "scope": dataset["purpose"], "model_calls":0, "passed":passed, "total":len(results),
              "engineering_success_rate": passed / max(len(results), 1), "results":results,
              "limitations":["受控同源模板回归，不是封存独立测试集", "不测自然语言语义、OCR 或真实用户收益", "不能替代风险 Macro-F1、模型消融和专业复核"]}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(f"Engineering regression: {passed}/{len(results)}; model calls: 0; report: {output}")
    return passed == len(results)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--data", type=Path, default=ROOT / "data/benchmarks/agent-regression.json")
    parser.add_argument("--output", type=Path, default=ROOT / ".data/agent-regression.json")
    args = parser.parse_args()
    raise SystemExit(0 if evaluate(args.data, args.output) else 1)
