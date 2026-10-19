import io
import json
from decimal import InvalidOperation
import pytest
from fastapi.testclient import TestClient
from tradepilot.api import create_app
from tradepilot.demos import demo_documents
from tradepilot.engine import extract, review, normalize, evidence_valid
from tradepilot.parsers import parse_document


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.delenv("TRADEPILOT_MODEL_URL", raising=False)
    app = create_app(data_dir=tmp_path, token="test-token")
    with TestClient(app, headers={"X-API-Key": "test-token"}) as c:
        yield c


def test_auth(client):
    assert client.get("/cases", headers={"X-API-Key": "wrong"}).status_code == 401
    assert client.get("/health").status_code == 200


def test_contract_difference_is_not_automatically_risk(client):
    case = client.post("/demo?variant=conflict").json()
    run = client.post(f"/cases/{case['id']}/runs").json()
    run = client.get(f"/cases/{case['id']}/runs/{run['id']}").json()
    assert run["status"] == "awaiting_review", run
    changes = [f for f in run["findings"] if f["kind"] == "difference"]
    risks = [f for f in run["findings"] if f["kind"] == "business_risk"]
    assert len(changes) == 3
    assert len(risks) == 2
    assert all(f["severity"] == "info" for f in changes)
    assert all(evidence_valid(e, case["documents"]) for f in run["findings"] for e in f["evidence"])


def test_approval_cannot_bypass_findings_or_limitations(client):
    case = client.post("/demo?variant=conflict").json()
    base = f"/cases/{case['id']}/runs"
    run = client.post(base).json()
    path = base + "/" + run["id"]
    run = client.get(path).json()
    assert client.post(path+"/approve", json={}).status_code == 409
    for finding in run["findings"]:
        assert client.post(path+"/decisions/"+finding["id"], json={"action":"accept", "reason":"已人工核对并接受条款变化"}).status_code == 200
    assert client.post(path+"/approve", json={}).status_code == 409
    approved = client.post(path+"/approve", json={"acknowledge_limitations":True,"reason":"已人工完成规则覆盖外的语义复核"})
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "approved"
    assert client.post(path+"/approve",json={}).json()["status"] == "approved"
    assert client.get(f"/cases/{case['id']}").json()["status"] == "archived"


def test_changes_invalidate_approvals_and_history_is_immutable(client):
    case = client.post("/demo?variant=normal").json()
    run = client.post(f"/cases/{case['id']}/runs").json()
    path = f"/cases/{case['id']}/runs/{run['id']}"
    original = client.get(path).json()
    contract = case["documents"][-1]
    fact = next(f for f in contract["facts"] if f["field"] == "total")
    result = client.patch(f"/cases/{case['id']}/documents/{contract['id']}/facts/{fact['id']}",json={"value":"95000", "reason":"核对原文后确认"})
    assert result.status_code == 200
    assert client.post(path+"/approve",json={"acknowledge_limitations":True,"reason":"test"}).status_code == 409
    old = client.get(path).json()
    old_total = next(f["value"] for f in old["input"]["documents"][-1]["facts"] if f["field"]=="total")
    assert old_total == "90000"
    second = client.post(f"/cases/{case['id']}/runs").json()
    assert second["id"] != run["id"]


def test_case_isolation_and_idempotency(client):
    a = client.post("/demo?variant=normal").json()
    b = client.post("/cases", json={"title":"Other"}).json()
    run = client.post(f"/cases/{a['id']}/runs").json()
    assert client.get(f"/cases/{b['id']}/runs/{run['id']}").status_code == 404
    again = client.post(f"/cases/{a['id']}/runs").json()
    assert again["id"] == run["id"]
    quote = demo_documents("normal")[1][2]
    repeated = client.post(f"/cases/{a['id']}/documents",data={"kind":"quote"},files={"file":("other-name.txt",quote.encode())})
    assert repeated.json()["id"] == a["documents"][1]["id"]
    assert client.put(f"/cases/{b['id']}/references",json=[a["documents"][0]["id"]]).status_code == 422


def test_exports(client):
    case = client.post("/demo?variant=normal").json()
    run = client.post(f"/cases/{case['id']}/runs").json()
    path = f"/cases/{case['id']}/runs/{run['id']}/export/"
    assert client.get(path+"pdf").content.startswith(b"%PDF")
    assert client.get(path+"docx").content.startswith(b"PK")
    assert client.get(path+"json").json()["run"]["id"] == run["id"]
    assert client.get(path+"exe").status_code == 422


def test_numeric_boundaries():
    assert normalize("total","90,000.00") == "90000"
    for value in ["NaN","Infinity","-1","101","1,23","1.000,00"]:
        with pytest.raises((ValueError,InvalidOperation)):
            normalize("deposit_percent",value)
    with pytest.raises(ValueError):
        normalize("shipment_date", "2027-02-30")


def test_office_parsing_and_formula_warning():
    from docx import Document
    from openpyxl import Workbook
    doc=Document();doc.add_paragraph("Seller: Acme");doc.add_table(rows=1,cols=1).cell(0,0).text="Total: USD 123"
    buf=io.BytesIO();doc.save(buf)
    blocks, warnings=parse_document("contract.docx",buf.getvalue())
    assert len(blocks)==2 and "table" in blocks[1]["location"]
    book=Workbook();book.active["A1"]="=1+2";buf=io.BytesIO();book.save(buf)
    blocks,warnings=parse_document("quote.xlsx",buf.getvalue())
    assert any("公式" in w for w in warnings)


def test_resource_guard_never_chooses_busy_device(monkeypatch):
    import importlib.util
    from pathlib import Path
    spec=importlib.util.spec_from_file_location("gpu_guard",Path(__file__).parents[1]/"scripts/gpu_guard.py")
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    def fake(fields,kind="gpu"):
        return [["GPU-test","81920","80000","0"]] if kind=="gpu" else [["GPU-test","42","0"]]
    monkeypatch.setattr(module,"query",fake)
    with pytest.raises(RuntimeError,match="计算进程"):
        module.inspect_gpu("GPU-test")
