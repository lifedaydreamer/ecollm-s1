import io
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from .schema import Block


def parse_document(name: str, content: bytes) -> tuple[list[dict], list[str]]:
    ext = Path(name).suffix.lower()
    blocks, warnings = [], []

    def add(text, page=None, location="", bbox=None):
        text = str(text).strip()
        if text:
            blocks.append(Block(id=f"b{len(blocks)+1}", text=text, page=page,
                                location=location, bbox=bbox).model_dump())

    if ext in {".docx", ".xlsx"}:
        try:
            with ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > 100 * 1024**2:
                    raise ValueError("解压后的文件超过 100MB 限制")
        except BadZipFile as exc:
            raise ValueError("损坏的 Office 文件") from exc
    if ext in {".txt", ".md", ".csv"}:
        for index, line in enumerate(content.decode("utf-8-sig").splitlines(), 1):
            add(line, location=f"line {index}")
    elif ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            if len(pdf.pages) > 100:
                raise ValueError("最多支持 100 页 PDF")
            for page_num, page in enumerate(pdf.pages, 1):
                lines = page.extract_text_lines()
                if not lines:
                    warnings.append(f"第 {page_num} 页缺少可提取文本，需要 OCR 后重新上传或人工补充")
                for line in lines:
                    add(line["text"], page_num, f"page {page_num}",
                        [line["x0"], line["top"], line["x1"], line["bottom"]])
    elif ext == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(io.BytesIO(content))
        for index, child in enumerate(doc.element.body, 1):
            if child.tag.endswith("}p"):
                add(Paragraph(child, doc).text, location=f"paragraph {index}")
            elif child.tag.endswith("}tbl"):
                for row_idx, row in enumerate(Table(child, doc).rows, 1):
                    add(" | ".join(c.text for c in row.cells), location=f"table {index} row {row_idx}")
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        formula = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        cached = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for sheet in formula:
            count = 0
            for row in sheet:
                parts = []
                for cell in row:
                    count += 1
                    if count > 100000:
                        raise ValueError("工作表超过 100000 个单元格限制")
                    if cell.value is None:
                        continue
                    value = cell.value
                    if cell.data_type == "f":
                        value = cached[sheet.title][cell.coordinate].value
                        if value is None:
                            warnings.append(f"{sheet.title}!{cell.coordinate} 的公式没有缓存结果")
                            continue
                    parts.append(f"{cell.coordinate}: {value}")
                if parts:
                    add(" | ".join(parts), location=f"{sheet.title}!row {row[0].row}")
        formula.close()
        cached.close()
    else:
        raise ValueError("支持 TXT、MD、CSV、PDF、DOCX、XLSX")
    if not blocks:
        warnings.append("未提取到正文，不能完成审查")
    return blocks, warnings
