from io import BytesIO
from html import escape


def report_text(case, run):
    lines = ["TradePilot 贸航 AI 审查报告", case["title"], f"审查编号 {run['id']} / 材料修订 {run['revision']}",
             f"状态 {run['status']} / 模型 {run.get('model', 'rules-only')}", ""]
    for limitation in run.get("limitations", []):
        lines.append("覆盖范围说明：" + limitation)
    for f in run.get("findings", []):
        lines += ["", f"[{f['severity']}] {f['title']}", f["explanation"], "建议：" + f["suggestion"]]
        for e in f["evidence"]:
            lines.append(f"证据 {e['document_id']} / {e['location']} / {e['block_id']}：{e['quote']}")
        for basis in f.get("basis", []):
            lines.append(f"依据或解释 {basis.get('title','')}：{basis.get('source','企业策略')} / {basis.get('version','')} {basis.get('text','')}")
        d = run.get("decisions", {}).get(f["id"])
        if d:
            lines.append(f"处理 {d['action']}：{d['reason']}")
    lines += ["", "英文商务邮件草稿", "Subject: Confirmation requested on proposed contract terms", 
              "Dear Partner,", "Please review the following points before we finalize the contract:"]
    lines += [f"- Please confirm the provision referenced in {f['evidence'][0]['location'] if f['evidence'] else 'the review report'} (review item {i+1})." for i, f in enumerate(run.get("findings", []))]
    lines += ["Please refer to the attached review for the original wording and proposed changes.", "Best regards"]
    return "\n".join(lines)


def export(case, run, fmt):
    text = report_text(case, run)
    if fmt == "txt":
        return text.encode(), "text/plain; charset=utf-8"
    if fmt == "pdf":
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        style = ParagraphStyle("body", fontName="STSong-Light", fontSize=10, leading=16, wordWrap="CJK", spaceAfter=5)
        buf = BytesIO()
        story = [Paragraph(escape(line) or "&#160;", style) for line in text.splitlines()]
        SimpleDocTemplate(buf, leftMargin=45, rightMargin=45, topMargin=40, bottomMargin=40).build(story)
        return buf.getvalue(), "application/pdf"
    if fmt == "docx":
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        doc.add_paragraph("TradePilot 合同审查与修改建议", "Title")
        doc.styles["Normal"].font.size = Pt(10)
        for line in text.splitlines()[1:]:
            doc.add_paragraph(line)
        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    raise ValueError("不支持的导出格式")
